import os
import re
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font


IMAGE_EXTENSIONS = [".jpg", ".jpeg", ".png", ".bmp", ".gif", ".webp", ".tif", ".tiff"]


def get_sheet_names(file_path):
    wb = load_workbook(file_path, read_only=True, data_only=True)
    return wb.sheetnames


def is_image_filename(value):
    if value is None:
        return False

    text = str(value).strip()
    lower_text = text.lower()

    return any(lower_text.endswith(ext) for ext in IMAGE_EXTENSIONS)


def parse_number(value):
    """
    尝试从数据列中提取数字。
    支持：
    123
    123.45
    -12.5
    123 mm
    数据=123.45
    """
    if value is None:
        return None

    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()

    match = re.search(r"-?\d+(?:\.\d+)?", text)

    if not match:
        return None

    try:
        return float(match.group())
    except ValueError:
        return None


def find_image_path(excel_file_path, label, data):
    excel_dir = os.path.dirname(excel_file_path)

    candidates = []

    if is_image_filename(label):
        candidates.append(str(label).strip())

    if is_image_filename(data):
        candidates.append(str(data).strip())

    for filename in candidates:
        image_path = os.path.join(excel_dir, filename)

        if os.path.exists(image_path):
            return image_path

        relative_path = os.path.join(
            excel_dir,
            filename.replace("/", os.sep).replace("\\", os.sep)
        )

        if os.path.exists(relative_path):
            return relative_path

    return ""


def read_excel_data(file_path, sheet_name):
    wb = load_workbook(file_path, data_only=True)
    ws = wb[sheet_name]

    rows = []
    n = 0

    file_name = os.path.basename(file_path)

    while True:
        base = 2 + 85 * n

        pairs = [
            (f"A{base}", f"C{base + 25}"),
            (f"J{base}", f"L{base + 25}"),
            (f"A{base + 42}", f"C{base + 67}"),
            (f"J{base + 42}", f"L{base + 67}"),
        ]

        if all(
            ws[label_cell].value is None and ws[data_cell].value is None
            for label_cell, data_cell in pairs
        ):
            break

        for label_cell, data_cell in pairs:
            label = ws[label_cell].value
            data = ws[data_cell].value

            if label is not None or data is not None:
                image_path = find_image_path(file_path, label, data)

                image_link_text = "打开图片" if image_path else ""

                rows.append([
                    file_name,
                    sheet_name,
                    label,
                    data,
                    image_link_text,
                    image_path
                ])

        n += 1

    return rows


def export_rows_to_excel(rows, save_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总结果"

    ws.append(["文件名", "Sheet", "标签", "数据", "图片链接", "图片完整路径"])

    for row in rows:
        file_name, sheet_name, label, data, image_link_text, image_path = row

        ws.append([
            file_name,
            sheet_name,
            label,
            data,
            image_link_text,
            image_path
        ])

        current_row = ws.max_row

        if image_path and os.path.exists(image_path):
            link_cell = ws.cell(row=current_row, column=5)
            link_cell.hyperlink = image_path
            link_cell.value = "打开图片"
            link_cell.font = Font(color="0000FF", underline="single")

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 80

    wb.save(save_path)


class ExcelReaderApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Excel 多文件多 Sheet 数据汇总工具")
        self.root.geometry("1150x760")

        self.file_records = []
        self.current_file_index = None

        # all_rows 保存全部读取结果
        # rows 保存当前筛选后的结果
        self.all_rows = []
        self.rows = []

        self.create_widgets()

    def create_widgets(self):
        top_frame = tk.Frame(self.root)
        top_frame.pack(fill=tk.X, padx=10, pady=10)

        choose_btn = tk.Button(
            top_frame,
            text="选择多个 Excel 文件",
            command=self.choose_files
        )
        choose_btn.pack(side=tk.LEFT)

        remove_btn = tk.Button(
            top_frame,
            text="移除选中文件",
            command=self.remove_selected_file
        )
        remove_btn.pack(side=tk.LEFT, padx=8)

        clear_btn = tk.Button(
            top_frame,
            text="清空文件列表",
            command=self.clear_files
        )
        clear_btn.pack(side=tk.LEFT, padx=8)

        read_btn = tk.Button(
            top_frame,
            text="读取并汇总",
            command=self.read_all_data
        )
        read_btn.pack(side=tk.LEFT, padx=8)

        open_btn = tk.Button(
            top_frame,
            text="打开选中图片",
            command=self.open_selected_image
        )
        open_btn.pack(side=tk.LEFT, padx=8)

        self.file_count_label = tk.Label(top_frame, text="已选择 0 个文件")
        self.file_count_label.pack(side=tk.LEFT, padx=20)

        file_frame = tk.LabelFrame(self.root, text="文件列表：点击某个文件后，可为它选择 Sheet")
        file_frame.pack(fill=tk.X, padx=10, pady=5)

        self.file_listbox = tk.Listbox(file_frame, height=7)
        self.file_listbox.pack(side=tk.LEFT, fill=tk.X, expand=True)

        file_scrollbar = tk.Scrollbar(file_frame, orient=tk.VERTICAL)
        file_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        self.file_listbox.config(yscrollcommand=file_scrollbar.set)
        file_scrollbar.config(command=self.file_listbox.yview)

        self.file_listbox.bind("<<ListboxSelect>>", self.on_file_selected)

        sheet_frame = tk.LabelFrame(self.root, text="当前选中文件的 Sheet 设置")
        sheet_frame.pack(fill=tk.X, padx=10, pady=5)

        self.current_file_label = tk.Label(sheet_frame, text="当前未选择文件", anchor="w")
        self.current_file_label.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)

        tk.Label(sheet_frame, text="Sheet：").pack(side=tk.LEFT)

        self.sheet_combo = ttk.Combobox(sheet_frame, state="readonly", width=30)
        self.sheet_combo.pack(side=tk.LEFT, padx=5)

        self.sheet_combo.bind("<<ComboboxSelected>>", self.on_sheet_selected)

        filter_frame = tk.LabelFrame(self.root, text="数据列大小筛选")
        filter_frame.pack(fill=tk.X, padx=10, pady=5)

        tk.Label(filter_frame, text="数据最小值：").pack(side=tk.LEFT, padx=5)

        self.min_value_entry = tk.Entry(filter_frame, width=12)
        self.min_value_entry.pack(side=tk.LEFT, padx=5)

        tk.Label(filter_frame, text="数据最大值：").pack(side=tk.LEFT, padx=5)

        self.max_value_entry = tk.Entry(filter_frame, width=12)
        self.max_value_entry.pack(side=tk.LEFT, padx=5)

        filter_btn = tk.Button(
            filter_frame,
            text="应用筛选",
            command=self.apply_data_filter
        )
        filter_btn.pack(side=tk.LEFT, padx=10)

        clear_filter_btn = tk.Button(
            filter_frame,
            text="清除筛选",
            command=self.clear_data_filter
        )
        clear_filter_btn.pack(side=tk.LEFT, padx=5)

        self.filter_status_label = tk.Label(filter_frame, text="当前未筛选")
        self.filter_status_label.pack(side=tk.LEFT, padx=15)

        table_frame = tk.Frame(self.root)
        table_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        self.tree = ttk.Treeview(
            table_frame,
            columns=("文件名", "Sheet", "标签", "数据", "图片链接", "图片完整路径"),
            show="headings"
        )

        self.tree.heading("文件名", text="文件名")
        self.tree.heading("Sheet", text="Sheet")
        self.tree.heading("标签", text="标签")
        self.tree.heading("数据", text="数据")
        self.tree.heading("图片链接", text="图片链接")
        self.tree.heading("图片完整路径", text="图片完整路径")

        self.tree.column("文件名", width=180)
        self.tree.column("Sheet", width=100)
        self.tree.column("标签", width=220)
        self.tree.column("数据", width=140)
        self.tree.column("图片链接", width=100)
        self.tree.column("图片完整路径", width=450)

        y_scrollbar = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=self.tree.yview)
        x_scrollbar = ttk.Scrollbar(table_frame, orient=tk.HORIZONTAL, command=self.tree.xview)

        self.tree.configure(
            yscrollcommand=y_scrollbar.set,
            xscrollcommand=x_scrollbar.set
        )

        self.tree.grid(row=0, column=0, sticky="nsew")
        y_scrollbar.grid(row=0, column=1, sticky="ns")
        x_scrollbar.grid(row=1, column=0, sticky="ew")

        table_frame.rowconfigure(0, weight=1)
        table_frame.columnconfigure(0, weight=1)

        self.tree.bind("<Double-1>", self.on_tree_double_click)

        bottom_frame = tk.Frame(self.root)
        bottom_frame.pack(fill=tk.X, padx=10, pady=10)

        self.status_label = tk.Label(bottom_frame, text="共 0 行")
        self.status_label.pack(side=tk.LEFT)

        hint_label = tk.Label(bottom_frame, text="提示：双击某一行可打开对应图片；导出的是当前筛选后的结果")
        hint_label.pack(side=tk.LEFT, padx=20)

        export_btn = tk.Button(
            bottom_frame,
            text="导出当前结果为 Excel",
            command=self.export_data
        )
        export_btn.pack(side=tk.RIGHT)

    def choose_files(self):
        file_paths = filedialog.askopenfilenames(
            title="选择多个 Excel 文件",
            filetypes=[
                ("Excel 文件", "*.xlsx *.xlsm"),
                ("所有文件", "*.*")
            ]
        )

        if not file_paths:
            return

        existing_paths = {record["path"] for record in self.file_records}
        failed_files = []

        for path in file_paths:
            if path in existing_paths:
                continue

            try:
                sheets = get_sheet_names(path)

                if not sheets:
                    failed_files.append(f"{os.path.basename(path)}：没有找到 Sheet")
                    continue

                record = {
                    "path": path,
                    "file_name": os.path.basename(path),
                    "sheets": sheets,
                    "selected_sheet": sheets[0],
                }

                self.file_records.append(record)

            except Exception as e:
                failed_files.append(f"{os.path.basename(path)}：{e}")

        self.refresh_file_list()

        if failed_files:
            messagebox.showwarning(
                "部分文件加载失败",
                "以下文件加载失败：\n\n" + "\n".join(failed_files)
            )

    def refresh_file_list(self):
        self.file_listbox.delete(0, tk.END)

        for index, record in enumerate(self.file_records, start=1):
            text = f"{index}. {record['file_name']}    当前 Sheet: {record['selected_sheet']}"
            self.file_listbox.insert(tk.END, text)

        self.file_count_label.config(text=f"已选择 {len(self.file_records)} 个文件")

    def on_file_selected(self, event=None):
        selection = self.file_listbox.curselection()

        if not selection:
            return

        index = selection[0]
        self.current_file_index = index

        record = self.file_records[index]

        self.current_file_label.config(text=record["path"])
        self.sheet_combo["values"] = record["sheets"]
        self.sheet_combo.set(record["selected_sheet"])

    def on_sheet_selected(self, event=None):
        if self.current_file_index is None:
            return

        selected_sheet = self.sheet_combo.get()

        if not selected_sheet:
            return

        self.file_records[self.current_file_index]["selected_sheet"] = selected_sheet

        current_index = self.current_file_index
        self.refresh_file_list()

        self.file_listbox.selection_set(current_index)
        self.file_listbox.activate(current_index)

    def remove_selected_file(self):
        selection = self.file_listbox.curselection()

        if not selection:
            messagebox.showwarning("提示", "请先选中要移除的文件。")
            return

        index = selection[0]
        del self.file_records[index]

        self.current_file_index = None
        self.current_file_label.config(text="当前未选择文件")
        self.sheet_combo["values"] = []
        self.sheet_combo.set("")

        self.refresh_file_list()

    def clear_files(self):
        self.file_records = []
        self.current_file_index = None
        self.all_rows = []
        self.rows = []

        self.file_listbox.delete(0, tk.END)

        for item in self.tree.get_children():
            self.tree.delete(item)

        self.current_file_label.config(text="当前未选择文件")
        self.sheet_combo["values"] = []
        self.sheet_combo.set("")

        self.min_value_entry.delete(0, tk.END)
        self.max_value_entry.delete(0, tk.END)

        self.file_count_label.config(text="已选择 0 个文件")
        self.status_label.config(text="共 0 行")
        self.filter_status_label.config(text="当前未筛选")

    def read_all_data(self):
        if not self.file_records:
            messagebox.showwarning("提示", "请先选择 Excel 文件。")
            return

        self.all_rows = []
        failed_files = []

        for record in self.file_records:
            file_path = record["path"]
            sheet_name = record["selected_sheet"]

            try:
                file_rows = read_excel_data(file_path, sheet_name)
                self.all_rows.extend(file_rows)

            except Exception as e:
                failed_files.append(
                    f"{record['file_name']} / {sheet_name}：{e}"
                )

        self.apply_data_filter(show_message=False)

        if failed_files:
            messagebox.showwarning(
                "部分文件读取失败",
                "以下文件读取失败：\n\n" + "\n".join(failed_files)
            )
        else:
            messagebox.showinfo(
                "完成",
                f"读取完成，共汇总 {len(self.file_records)} 个文件，原始 {len(self.all_rows)} 行数据。"
            )

    def apply_data_filter(self, show_message=True):
        min_text = self.min_value_entry.get().strip()
        max_text = self.max_value_entry.get().strip()

        try:
            min_value = float(min_text) if min_text else None
        except ValueError:
            messagebox.showerror("错误", "数据最小值必须是数字。")
            return

        try:
            max_value = float(max_text) if max_text else None
        except ValueError:
            messagebox.showerror("错误", "数据最大值必须是数字。")
            return

        if min_value is not None and max_value is not None and min_value > max_value:
            messagebox.showerror("错误", "最小值不能大于最大值。")
            return

        if min_value is None and max_value is None:
            self.rows = list(self.all_rows)
            self.filter_status_label.config(text="当前未筛选")
        else:
            filtered_rows = []

            for row in self.all_rows:
                data_value = row[3]
                number = parse_number(data_value)

                if number is None:
                    continue

                if min_value is not None and number < min_value:
                    continue

                if max_value is not None and number > max_value:
                    continue

                filtered_rows.append(row)

            self.rows = filtered_rows

            if min_value is not None and max_value is not None:
                text = f"筛选：{min_value} ≤ 数据 ≤ {max_value}"
            elif min_value is not None:
                text = f"筛选：数据 ≥ {min_value}"
            else:
                text = f"筛选：数据 ≤ {max_value}"

            self.filter_status_label.config(text=text)

        self.refresh_table()

        if show_message:
            messagebox.showinfo(
                "筛选完成",
                f"原始 {len(self.all_rows)} 行，当前显示 {len(self.rows)} 行。"
            )

    def clear_data_filter(self):
        self.min_value_entry.delete(0, tk.END)
        self.max_value_entry.delete(0, tk.END)

        self.rows = list(self.all_rows)

        self.filter_status_label.config(text="当前未筛选")
        self.refresh_table()

    def refresh_table(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        for row in self.rows:
            self.tree.insert("", tk.END, values=row)

        self.status_label.config(
            text=f"当前 {len(self.rows)} 行 / 原始 {len(self.all_rows)} 行"
        )

    def get_selected_image_path(self):
        selected_items = self.tree.selection()

        if not selected_items:
            return ""

        item_id = selected_items[0]
        values = self.tree.item(item_id, "values")

        if len(values) < 6:
            return ""

        image_path = values[5]

        return image_path

    def open_selected_image(self):
        image_path = self.get_selected_image_path()

        if not image_path:
            messagebox.showwarning("提示", "当前选中行没有找到图片链接。")
            return

        if not os.path.exists(image_path):
            messagebox.showerror("错误", f"图片文件不存在：\n{image_path}")
            return

        try:
            os.startfile(image_path)
        except Exception as e:
            messagebox.showerror("错误", f"打开图片失败：\n{e}")

    def on_tree_double_click(self, event=None):
        self.open_selected_image()

    def export_data(self):
        if not self.rows:
            messagebox.showwarning("提示", "当前没有可导出的数据。")
            return

        save_path = filedialog.asksaveasfilename(
            title="保存汇总结果",
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")]
        )

        if not save_path:
            return

        try:
            export_rows_to_excel(self.rows, save_path)
            messagebox.showinfo("成功", f"已导出当前结果到：\n{save_path}")

        except Exception as e:
            messagebox.showerror("错误", f"导出失败：\n{e}")


if __name__ == "__main__":
    root = tk.Tk()
    app = ExcelReaderApp(root)
    root.mainloop()